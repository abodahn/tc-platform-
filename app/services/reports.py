"""
TC Platform — Reports & Exports engine.

A registry of real, filterable reports that run live SQL against the platform's
own data (procurement, maintenance, production, and platform metadata) and
export to CSV / Excel / PDF. Each report is a declarative spec; filters are
whitelisted columns bound as parameters (no SQL injection), and every report is
gated by a permission so users only see what their role allows.
"""
from __future__ import annotations

import io

from app.db import get_db

_PR_STATUS = ["draft", "pending", "approved", "rejected", "po_issued",
              "partially_received", "received", "closed"]
_PRIORITY = ["low", "medium", "high", "critical"]
_SEVERITY = ["info", "warning", "critical"]

# Reusable filter builders --------------------------------------------------
def _f(name, label, col, type="text", op="like", options=None):
    return {"name": name, "label": label, "col": col, "type": type, "op": op, "options": options}

def _date_range(col="created_at"):
    return [_f("from", "From", col, "date", ">="), _f("to", "To", col, "date", "<=")]


# Report registry -----------------------------------------------------------
# Each: key, group, label, desc, perm, icon, sql (with {where}), columns[(key,header)], filters, base_where
REPORTS = [
    # ---------------- Procurement ----------------
    dict(key="proc_prs", group="Procurement", label="Purchase Requests",
         desc="Every purchase request with status, vendor, value and payment.",
         perm="proc_view", icon="i-cart",
         sql="SELECT pr_no,title,department,requester_name AS requester,vendor,total,"
             "currency,status,payment_status,created_at FROM pr_requests {where} ORDER BY id DESC",
         columns=[("pr_no", "PR No"), ("title", "Title"), ("department", "Department"),
                  ("requester", "Requester"), ("vendor", "Vendor"), ("total", "Total"),
                  ("currency", "Cur"), ("status", "Status"), ("payment_status", "Payment"),
                  ("created_at", "Created")],
         filters=[_f("status", "Status", "status", "select", "=", _PR_STATUS),
                  _f("department", "Department", "department"),
                  _f("vendor", "Vendor", "vendor")] + _date_range()),
    dict(key="proc_spend_dept", group="Procurement", label="Spend by Department",
         desc="Total committed spend and request count per department.",
         perm="proc_view", icon="i-chart",
         sql="SELECT department,COUNT(*) AS requests,SUM(total) AS total FROM pr_requests "
             "{where} GROUP BY department ORDER BY total DESC",
         columns=[("department", "Department"), ("requests", "Requests"), ("total", "Total spend")],
         filters=[_f("status", "Status", "status", "select", "=", _PR_STATUS)] + _date_range()),
    dict(key="proc_invoices", group="Procurement", label="Vendor Invoices",
         desc="Invoices recorded against purchase orders (for 3-way match).",
         perm="proc_view", icon="i-report",
         sql="SELECT i.invoice_no,p.pr_no,i.invoice_date,i.amount,i.tax,i.currency,i.status "
             "FROM pr_invoices i LEFT JOIN pr_requests p ON p.id=i.pr_id {where} ORDER BY i.id DESC",
         columns=[("invoice_no", "Invoice"), ("pr_no", "PR No"), ("invoice_date", "Date"),
                  ("amount", "Amount"), ("tax", "Tax"), ("currency", "Cur"), ("status", "Status")],
         filters=[_f("status", "Status", "i.status", "select", "=",
                     ["received", "matched", "disputed", "paid"])]),
    dict(key="proc_payments", group="Procurement", label="Payments",
         desc="Payments made against purchase requests and invoices.",
         perm="proc_view", icon="i-wallet",
         sql="SELECT y.paid_at,p.pr_no,y.amount,y.currency,y.method,y.reference "
             "FROM pr_payments y LEFT JOIN pr_requests p ON p.id=y.pr_id {where} ORDER BY y.id DESC",
         columns=[("paid_at", "Paid on"), ("pr_no", "PR No"), ("amount", "Amount"),
                  ("currency", "Cur"), ("method", "Method"), ("reference", "Reference")],
         filters=[_f("method", "Method", "y.method")]),
    dict(key="proc_budgets", group="Procurement", label="Department Budgets",
         desc="Annual department budgets on record.",
         perm="proc_view", icon="i-wallet",
         sql="SELECT department,period,amount,currency FROM proc_budgets {where} "
             "ORDER BY period DESC, department",
         columns=[("department", "Department"), ("period", "Year"), ("amount", "Budget"),
                  ("currency", "Cur")],
         filters=[_f("department", "Department", "department")]),
    dict(key="proc_vendors", group="Procurement", label="Vendors",
         desc="Supplier master with contact, category and rating.",
         perm="proc_view", icon="i-boxes",
         sql="SELECT name,contact_person,phone,email,category,rating,is_active "
             "FROM proc_vendors {where} ORDER BY name",
         columns=[("name", "Vendor"), ("contact_person", "Contact"), ("phone", "Phone"),
                  ("email", "Email"), ("category", "Category"), ("rating", "Rating"),
                  ("is_active", "Active")],
         filters=[_f("category", "Category", "category")]),

    # ---------------- Maintenance ----------------
    dict(key="mnt_tickets", group="Maintenance", label="Maintenance Tickets",
         desc="Work orders with priority, status, downtime and cost.",
         perm="maint_view", icon="i-report",
         sql="SELECT ticket_no,machine_code,department,issue_category,priority,status,"
             "assigned_to,total_downtime_min,cost,created_at,closed_at FROM mnt_tickets "
             "{where} ORDER BY id DESC",
         columns=[("ticket_no", "Ticket"), ("machine_code", "Machine"), ("department", "Department"),
                  ("issue_category", "Category"), ("priority", "Priority"), ("status", "Status"),
                  ("assigned_to", "Assigned"), ("total_downtime_min", "Downtime (min)"),
                  ("cost", "Cost"), ("created_at", "Created"), ("closed_at", "Closed")],
         filters=[_f("status", "Status", "status"),
                  _f("priority", "Priority", "priority", "select", "=", _PRIORITY),
                  _f("department", "Department", "department")] + _date_range()),
    dict(key="mnt_machines", group="Maintenance", label="Machine Register",
         desc="Machines with criticality, status, breakdowns and next PM.",
         perm="maint_view", icon="i-factory",
         sql="SELECT code,name,type,department,criticality,status,next_pm_date,breakdowns,"
             "total_downtime_min,cost_to_date FROM mnt_machines {where} ORDER BY code",
         columns=[("code", "Code"), ("name", "Name"), ("type", "Type"), ("department", "Department"),
                  ("criticality", "Criticality"), ("status", "Status"), ("next_pm_date", "Next PM"),
                  ("breakdowns", "Breakdowns"), ("total_downtime_min", "Downtime (min)"),
                  ("cost_to_date", "Cost to date")],
         filters=[_f("department", "Department", "department"),
                  _f("status", "Status", "status"),
                  _f("criticality", "Criticality", "criticality", "select", "=", _PRIORITY)]),
    dict(key="mnt_spares_low", group="Maintenance", label="Spare Parts — Low Stock",
         desc="Parts at or below reorder level — reorder now.",
         perm="maint_view", icon="i-boxes",
         sql="SELECT code,name,category,stock_qty,reorder_level,min_level,avg_cost,warehouse,"
             "criticality,vendor FROM mnt_spare_parts {where} ORDER BY (stock_qty-reorder_level)",
         columns=[("code", "Code"), ("name", "Part"), ("category", "Category"),
                  ("stock_qty", "In stock"), ("reorder_level", "Reorder at"), ("min_level", "Min"),
                  ("avg_cost", "Avg cost"), ("warehouse", "Warehouse"),
                  ("criticality", "Criticality"), ("vendor", "Vendor")],
         base_where=["stock_qty <= reorder_level", "is_active = 1"],
         filters=[_f("category", "Category", "category"),
                  _f("criticality", "Criticality", "criticality", "select", "=", _PRIORITY)]),
    dict(key="mnt_spares", group="Maintenance", label="Spare Parts Inventory",
         desc="Full spare-parts inventory with stock and cost.",
         perm="maint_view", icon="i-boxes",
         sql="SELECT code,name,category,uom,stock_qty,reorder_level,avg_cost,warehouse,criticality "
             "FROM mnt_spare_parts {where} ORDER BY code",
         columns=[("code", "Code"), ("name", "Part"), ("category", "Category"), ("uom", "UoM"),
                  ("stock_qty", "In stock"), ("reorder_level", "Reorder at"), ("avg_cost", "Avg cost"),
                  ("warehouse", "Warehouse"), ("criticality", "Criticality")],
         filters=[_f("category", "Category", "category")]),

    # ---------------- Production ----------------
    dict(key="prod_lines", group="Production", label="Production Lines",
         desc="Line status, shift and output.",
         perm="view_reports", icon="i-factory",
         sql="SELECT name,area,status,shift,target_output,actual_output,operators,updated_at "
             "FROM production_lines {where} ORDER BY area,name",
         columns=[("name", "Line"), ("area", "Area"), ("status", "Status"), ("shift", "Shift"),
                  ("target_output", "Target"), ("actual_output", "Actual"),
                  ("operators", "Operators"), ("updated_at", "Updated")],
         filters=[_f("status", "Status", "status"), _f("area", "Area", "area")]),
    dict(key="prod_downtime", group="Production", label="Production Downtime",
         desc="Downtime events by line and category.",
         perm="view_reports", icon="i-clock",
         sql="SELECT l.name AS line,d.reason,d.category,d.minutes,d.occurred_at "
             "FROM production_downtime d LEFT JOIN production_lines l ON l.id=d.line_id "
             "{where} ORDER BY d.id DESC",
         columns=[("line", "Line"), ("reason", "Reason"), ("category", "Category"),
                  ("minutes", "Minutes"), ("occurred_at", "Occurred")],
         filters=[_f("category", "Category", "d.category")]),
    dict(key="prod_quality", group="Production", label="Production Quality",
         desc="Quality issues by line, severity and status.",
         perm="view_reports", icon="i-alert",
         sql="SELECT l.name AS line,q.issue,q.severity,q.quantity,q.status,q.created_at "
             "FROM production_quality q LEFT JOIN production_lines l ON l.id=q.line_id "
             "{where} ORDER BY q.id DESC",
         columns=[("line", "Line"), ("issue", "Issue"), ("severity", "Severity"),
                  ("quantity", "Qty"), ("status", "Status"), ("created_at", "Created")],
         filters=[_f("status", "Status", "q.status")]),

    # ---------------- Platform (admin) ----------------
    dict(key="audit", group="Platform", label="Audit Log",
         desc="Every recorded user action.",
         perm="access_admin", icon="i-report",
         sql="SELECT id,username,action,detail,ip,created_at FROM audit_logs {where} ORDER BY id DESC",
         columns=[("id", "#"), ("username", "User"), ("action", "Action"), ("detail", "Detail"),
                  ("ip", "IP"), ("created_at", "When")],
         filters=[_f("username", "User", "username"), _f("action", "Action", "action")] + _date_range()),
    dict(key="notifications", group="Platform", label="Notifications",
         desc="Aggregated alerts across all systems.",
         perm="access_admin", icon="i-bell",
         sql="SELECT id,severity,module,title,message,is_read,created_at FROM notifications "
             "{where} ORDER BY id DESC",
         columns=[("id", "#"), ("severity", "Severity"), ("module", "Module"), ("title", "Title"),
                  ("message", "Message"), ("is_read", "Read"), ("created_at", "When")],
         filters=[_f("severity", "Severity", "severity", "select", "=", _SEVERITY),
                  _f("module", "Module", "module")]),
    dict(key="users", group="Platform", label="Users & Roles",
         desc="User accounts, roles and status.",
         perm="access_admin", icon="i-users",
         sql="SELECT username,full_name,email,role,is_active,created_at FROM users {where} "
             "ORDER BY username",
         columns=[("username", "Username"), ("full_name", "Name"), ("email", "Email"),
                  ("role", "Role"), ("is_active", "Active"), ("created_at", "Created")],
         filters=[_f("role", "Role", "role")]),
    dict(key="systems", group="Platform", label="Systems Registry",
         desc="Integrated systems and internal modules.",
         perm="access_admin", icon="i-grid",
         sql="SELECT key,name_en AS name,category,base_url,owner,criticality,is_integrated,enabled "
             "FROM systems {where} ORDER BY sort_order",
         columns=[("key", "Key"), ("name", "Name"), ("category", "Category"), ("base_url", "URL"),
                  ("owner", "Owner"), ("criticality", "Criticality"), ("is_integrated", "Integrated"),
                  ("enabled", "Enabled")],
         filters=[_f("category", "Category", "category")]),
]

_BY_KEY = {r["key"]: r for r in REPORTS}
GROUP_ORDER = ["Procurement", "Maintenance", "Production", "Platform"]


def get(key):
    return _BY_KEY.get(key)


def catalog(can):
    """Group reports the user may view. `can` is a callable(perm) -> bool."""
    out = []
    for g in GROUP_ORDER:
        items = [r for r in REPORTS if r["group"] == g and can(r["perm"])]
        if items:
            out.append({"group": g, "reports": items})
    return out


def _build_where(spec, args):
    clauses = list(spec.get("base_where", []))
    params = []
    for f in spec.get("filters", []):
        v = (args.get(f["name"]) or "").strip()
        if not v:
            continue
        op = f.get("op", "like")
        col = f["col"]
        if op == "like":
            clauses.append(f"{col} LIKE ?")
            params.append(f"%{v}%")
        else:
            if f.get("type") == "date" and op == "<=" and len(v) == 10:
                v = v + " 23:59:59"
            clauses.append(f"{col} {op} ?")
            params.append(v)
    where = ("WHERE " + " AND ".join(clauses)) if clauses else ""
    return where, params


def run(key, args, limit=None):
    spec = _BY_KEY.get(key)
    if not spec:
        return None
    where, params = _build_where(spec, args)
    sql = spec["sql"].replace("{where}", where)
    if limit:
        sql += f" LIMIT {int(limit)}"
    conn = get_db()
    try:
        rows = conn.execute(sql, tuple(params)).fetchall()
    finally:
        conn.close()
    return {"spec": spec, "rows": rows, "count": len(rows)}


# --- exporters --------------------------------------------------------------
def _cell(v):
    """Neutralise spreadsheet formula injection on string cells."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def _val(row, key):
    try:
        v = row[key]
    except Exception:
        v = None
    return "" if v is None else v


def export(spec, rows, fmt):
    """Return (payload_bytes, mimetype, extension) for the given format."""
    cols = spec["columns"]
    headers = [h for _k, h in cols]
    keys = [k for k, _h in cols]
    if fmt == "csv":
        import csv
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(headers)
        for r in rows:
            w.writerow([_cell(_val(r, k)) for k in keys])
        return buf.getvalue().encode("utf-8-sig"), "text/csv", "csv"

    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = spec["label"][:31]
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="0B1222")
        for r in rows:
            ws.append([_cell(_val(r, k)) for k in keys])
        for i, _h in enumerate(headers, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 18
        ws.freeze_panes = "A2"
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue(), \
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"

    if fmt == "pdf":
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.units import cm
        from reportlab.lib import colors
        from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
                                         Paragraph, Spacer)
        from reportlab.lib.styles import getSampleStyleSheet
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                                leftMargin=1.2 * cm, rightMargin=1.2 * cm,
                                topMargin=1.2 * cm, bottomMargin=1.2 * cm,
                                title=f"TC Platform — {spec['label']}")
        styles = getSampleStyleSheet()
        story = [Paragraph(f"<b>{spec['label']}</b>", styles["Title"]),
                 Paragraph(f"TC Platform &middot; {len(rows)} rows", styles["Normal"]),
                 Spacer(1, 0.4 * cm)]

        def _trim(v):
            s = str(v)
            return s if len(s) <= 42 else s[:40] + "…"
        data = [headers] + [[_trim(_val(r, k)) for k in keys] for r in rows[:2500]]
        tbl = Table(data, repeatRows=1)
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B1222")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9DEE7")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F6FA")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(tbl)
        doc.build(story)
        return buf.getvalue(), "application/pdf", "pdf"

    raise ValueError("bad format")
