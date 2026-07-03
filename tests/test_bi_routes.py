"""Route + store + jobs tests for BI, against an isolated temp DB so the real
platform.db is never touched."""
import io

import pytest

from _support import login_admin, get_csrf


@pytest.fixture()
def bi(tmp_path, monkeypatch):
    from config import Config
    monkeypatch.setattr(Config, "DB_PATH", tmp_path / "bi.db")
    monkeypatch.setattr(Config, "DATABASE_URL", "")
    import app.services.bi.store as store
    store._CACHE.clear()
    from app import create_app
    a = create_app()
    a.config["TESTING"] = True
    c = a.test_client()
    login_admin(c)
    return a, c


def _upload(c, text=b"City,Sales,Date\nNY,100,2026-01-01\nLA,200,2026-02-01\nSF,150,2026-03-01\nNY,120,2026-04-01\n"):
    tok = get_csrf(c)
    return c.post("/bi/upload", data={"file": (io.BytesIO(text), "s.csv")},
                  content_type="multipart/form-data", headers={"X-CSRF-Token": tok}), tok


def test_upload_and_dataset_page(bi):
    a, c = bi
    r, tok = _upload(c)
    assert r.status_code == 200
    dsid = r.get_json()["dataset_id"]
    assert c.get(f"/bi/dataset/{dsid}").status_code == 200


def test_upload_rejects_bad_type(bi):
    a, c = bi
    tok = get_csrf(c)
    r = c.post("/bi/upload", data={"file": (io.BytesIO(b"x"), "x.exe")},
               content_type="multipart/form-data", headers={"X-CSRF-Token": tok})
    assert r.status_code == 400


def test_sample_and_apis(bi):
    a, c = bi
    tok = get_csrf(c)
    r = c.post("/bi/sample", headers={"X-CSRF-Token": tok})
    dsid = r.get_json()["dataset_id"]
    # auto dashboard
    meta = c.get(f"/bi/api/dataset/{dsid}").get_json()
    assert meta["dashboard"]["kpis"] and meta["dashboard"]["charts"]
    # insights localised
    for lang in ("en", "ar", "tr"):
        ins = c.get(f"/bi/api/dataset/{dsid}/insights?lang={lang}").get_json()["insights"]
        assert ins
    # a chart
    ch = meta["dashboard"]["charts"][0]
    cd = c.post(f"/bi/api/dataset/{dsid}/chart", json={"spec": ch, "filters": []},
                headers={"X-CSRF-Token": tok}).get_json()
    assert "kind" in cd
    # kpis under a filter
    k = c.post(f"/bi/api/dataset/{dsid}/kpis",
               json={"kpis": meta["dashboard"]["kpis"], "filters": [{"column": "Region", "op": "=", "value": "North"}]},
               headers={"X-CSRF-Token": tok}).get_json()
    assert "kpis" in k
    # smart search
    q = c.post(f"/bi/api/dataset/{dsid}/query", json={"q": "top 3 product by revenue"},
               headers={"X-CSRF-Token": tok}).get_json()
    assert q["matched"] is True


def test_save_pin_and_view_dashboard(bi):
    a, c = bi
    r, tok = _upload(c)
    dsid = r.get_json()["dataset_id"]
    meta = c.get(f"/bi/api/dataset/{dsid}").get_json()
    saved = c.post("/bi/dashboard/save", json={"dataset_id": dsid, "name": "D1", "spec": meta["dashboard"]},
                   headers={"X-CSRF-Token": tok}).get_json()
    did = saved["id"]
    assert c.get(f"/bi/dashboard/{did}").status_code == 200
    assert c.post(f"/bi/dashboard/{did}/pin", json={"pinned": True}, headers={"X-CSRF-Token": tok}).get_json()["ok"]


def test_exports(bi):
    a, c = bi
    r, tok = _upload(c)
    dsid = r.get_json()["dataset_id"]
    xr = c.get(f"/bi/export/{dsid}.xlsx")
    assert xr.status_code == 200 and len(xr.data) > 500
    pr = c.get(f"/bi/export/{dsid}.pdf")
    assert pr.status_code == 200 and pr.data[:4] == b"%PDF"


def test_alert_fires_into_notifications(bi):
    a, c = bi
    r, tok = _upload(c)
    dsid = r.get_json()["dataset_id"]
    c.post("/bi/alert", json={"dataset_id": dsid, "column_name": "Sales", "agg": "sum",
                              "op": ">", "threshold": 1}, headers={"X-CSRF-Token": tok})
    from app.services.bi import jobs
    with a.app_context():
        n = jobs.evaluate_alerts()
        assert n == 1
        from app.db import get_db
        conn = get_db()
        row = conn.execute("SELECT COUNT(*) c FROM notifications WHERE module='bi'").fetchone()
        conn.close()
        assert row["c"] >= 1
        # second run: no duplicate (still breach, no fresh transition)
        assert jobs.evaluate_alerts() == 0


def test_jobs_run_requires_auth(bi):
    a, c = bi
    # admin session -> allowed
    r = c.post("/bi/jobs/run", headers={"X-CSRF-Token": get_csrf(c)})
    assert r.status_code == 200
    # anonymous -> 403
    anon = a.test_client()
    r2 = anon.post("/bi/jobs/run")
    assert r2.status_code in (400, 403)  # 400 if CSRF gate hits first, 403 otherwise


def test_export_xlsx_sanitizes_formula_injection(bi):
    a, c = bi
    tok = get_csrf(c)
    payload = b"Name,Note\nAcme,=HYPERLINK(1)\nBeta,+danger\n"
    r = c.post("/bi/upload", data={"file": (io.BytesIO(payload), "f.csv")},
               content_type="multipart/form-data", headers={"X-CSRF-Token": tok})
    dsid = r.get_json()["dataset_id"]
    xr = c.get(f"/bi/export/{dsid}.xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xr.data))
    vals = [cell.value for row in wb["Data"].iter_rows() for cell in row]
    # formula-triggering cells are neutralised with a leading quote
    assert "'=HYPERLINK(1)" in vals
    assert "'+danger" in vals


def test_store_owner_scoping(bi):
    a, c = bi
    with a.app_context():
        from app.services.bi import store
        prof = {"columns": [], "n_rows": 1}
        qual = {"score": 100, "issues": []}
        store.save_dataset("A", "a.csv", "upload", ["x"], [[1]], prof, qual, "alice")
        store.save_dataset("B", "b.csv", "upload", ["x"], [[1]], prof, qual, "bob")
        alice = [d["name"] for d in store.list_datasets(owner="alice")]
        everyone = [d["name"] for d in store.list_datasets()]
        assert "A" in alice and "B" not in alice        # scoped
        assert "A" in everyone and "B" in everyone       # admin/all view
