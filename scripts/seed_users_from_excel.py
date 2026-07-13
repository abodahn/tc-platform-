"""
Bulk-create platform user accounts from an Excel/CSV export (e.g. the Exchange
"ALL TC" address list: Display Name / Alias / Primary SMTP Address).

Every user is created with the **itsm_user** role by default — a minimal account
that can sign in and open the IT Service Desk, with no procurement / probation /
admin access. Idempotent: existing usernames are skipped (never overwritten).

Credentials: unless TC_SEED_PASSWORD is set (one shared temp password for all),
a strong random password is generated per user and written to a CSV **next to
the input file** (outside this repo — employee names/emails are never committed).

Usage:
    python scripts/seed_users_from_excel.py "C:\\path\\to\\ALL TC.xlsx"
    python scripts/seed_users_from_excel.py            # defaults to Downloads\\ALL TC.xlsx
Options via env:
    USERS_ROLE=itsm_user        role to assign
    TC_SEED_PASSWORD=...         one shared password instead of per-user random

Works on SQLite (local/on-prem) and PostgreSQL (Render) via app.db.get_db().
"""
import os
import sys
import csv
import secrets
import string
from datetime import datetime, timezone

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
os.chdir(ROOT)

from werkzeug.security import generate_password_hash  # noqa: E402
from app import create_app  # noqa: E402
from app.db import get_db  # noqa: E402

ROLE = os.environ.get("USERS_ROLE", "itsm_user").strip() or "itsm_user"

# Header aliases (case-insensitive) -> field
_HDR = {
    "username": ("alias", "username", "user name", "login", "samaccountname", "user id"),
    "full_name": ("display name", "name", "full name", "displayname"),
    "email": ("primary smtp address", "email", "e-mail", "mail", "smtp", "email address"),
}


def _norm(s):
    return ("" if s is None else str(s)).strip()


def _read_rows(path):
    """Return list of {username, full_name, email} from a .xlsx/.csv, tolerant of
    a header row that carries the known column names."""
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        grids = [list(ws.iter_rows(values_only=True)) for ws in wb.worksheets]
    else:
        import io
        with open(path, "r", encoding="utf-8-sig", errors="replace") as f:
            grids = [[tuple(r) for r in csv.reader(f)]]
    out = []
    for grid in grids:
        if not grid:
            continue
        # find the header row within the first few rows
        hdr_i, cols = None, {}
        for i, row in enumerate(grid[:6]):
            cells = [_norm(c).lower() for c in row]
            found = {}
            for field, names in _HDR.items():
                for j, c in enumerate(cells):
                    if c in names:
                        found[field] = j
                        break
            if "email" in found or "username" in found:
                hdr_i, cols = i, found
                break
        if hdr_i is None:
            continue
        for row in grid[hdr_i + 1:]:
            def cell(field):
                j = cols.get(field)
                return _norm(row[j]) if (j is not None and j < len(row)) else ""
            email = cell("email")
            username = cell("username") or (email.split("@")[0] if email else "")
            if not username:
                continue
            out.append({"username": username.lower(), "full_name": cell("full_name"),
                        "email": email})
    return out


def _password():
    shared = os.environ.get("TC_SEED_PASSWORD")
    if shared:
        return lambda: shared, False
    alphabet = string.ascii_letters + string.digits + "!@#$%*?"
    return (lambda: "".join(secrets.choice(alphabet) for _ in range(12))), True


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        os.path.expanduser("~"), "Downloads", "ALL TC.xlsx")
    if not os.path.exists(path):
        print("Input file not found:", path)
        sys.exit(1)
    rows = _read_rows(path)
    if not rows:
        print("No user rows recognised in", path)
        sys.exit(1)

    gen_pw, random_pw = _password()
    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
    app = create_app()
    created, skipped, errors, creds = 0, 0, 0, []
    with app.app_context():
        conn = get_db()
        try:
            conn.execute("SELECT account_status FROM users LIMIT 1").fetchone()
            has_status = True
        except Exception:
            try:
                conn.rollback()
            except Exception:
                pass
            has_status = False
        try:
            seen = set()
            for r in rows:
                u = r["username"]
                if u in seen:
                    continue
                seen.add(u)
                try:
                    exists = conn.execute("SELECT id FROM users WHERE username=?", (u,)).fetchone()
                    if exists:
                        skipped += 1
                        continue
                    pw = gen_pw()
                    ph = generate_password_hash(pw)
                    if has_status:
                        conn.execute(
                            "INSERT INTO users (username,password_hash,full_name,email,role,is_active,"
                            "account_status,created_at) VALUES (?,?,?,?,?,1,'active',?)",
                            (u, ph, r["full_name"] or u, r["email"], ROLE, now))
                    else:
                        conn.execute(
                            "INSERT INTO users (username,password_hash,full_name,email,role,is_active,"
                            "created_at) VALUES (?,?,?,?,?,1,?)",
                            (u, ph, r["full_name"] or u, r["email"], ROLE, now))
                    created += 1
                    if random_pw:
                        creds.append((u, r["full_name"], r["email"], pw))
                except Exception as exc:  # noqa: BLE001
                    errors += 1
                    print("  error on", u, "->", str(exc)[:100])
            conn.commit()
        finally:
            conn.close()

    print("=" * 60)
    print(" Users seeded from:", os.path.basename(path))
    print("  role     :", ROLE)
    print("  in file  :", len(rows))
    print("  created  :", created)
    print("  skipped  :", skipped, "(already existed)")
    print("  errors   :", errors)
    if random_pw and creds:
        out_csv = os.path.join(os.path.dirname(os.path.abspath(path)),
                               "itsm_users_credentials.csv")
        with open(out_csv, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(["username", "full_name", "email", "temp_password"])
            w.writerows(creds)
        print("\n  Credentials for the", len(creds), "new users written to:")
        print("   ", out_csv)
        print("  (distribute privately; users change the password after first login)")
    elif not random_pw:
        print("\n  Password: the shared TC_SEED_PASSWORD you set (change after first login).")
    print("=" * 60)


if __name__ == "__main__":
    main()
